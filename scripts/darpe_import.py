import json
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl


WORKBOOK_PATH = Path("scripts/darpe_import_source.xlsx")
OUTPUT_DIR = Path("scripts/darpe_import_output")
IMPORT_MARKER = "IMPORTADO_DARPE_PLANILHA_2026_04_19"

MONTHS = {
    "JANEIRO": 1,
    "FEVEREIRO": 2,
    "MARCO": 3,
    "MARCOO": 3,
    "ABRIL": 4,
    "MAIO": 5,
    "JUNHO": 6,
    "JULHO": 7,
    "AGOSTO": 8,
    "SETEMBRO": 9,
    "OUTUBRO": 10,
    "NOVEMBRO": 11,
    "DEZEMBRO": 12,
}

CITY_ALIASES = {
    "CAUCAIA DO ALTO": "CAUCAIA DO ALTO",
    "VGP": "VARGEM GRANDE PAULISTA",
    "VARGEM GRANDE PAULISTA": "VARGEM GRANDE PAULISTA",
    "COTIA": "COTIA",
    "ITAPEVI": "ITAPEVI",
    "JANDIRA": "JANDIRA",
    "CARAPICUIBA": "CARAPICUIBA",
    "OSASCO": "OSASCO",
    "BARUERI": "BARUERI",
    "SAO ROQUE": "SAO ROQUE",
}


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).replace("\xa0", " ").strip()
    text = re.sub(r"\s+", " ", text)
    return text


def strip_accents(value):
    return "".join(
        char
        for char in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(char)
    )


def slugify(value):
    base = strip_accents(clean_text(value)).upper()
    base = re.sub(r"\([^)]*\)", " ", base)
    base = re.sub(r"[^A-Z0-9]+", " ", base)
    base = re.sub(r"\b(ASSOCIACAO|DE|DA|DO|DAS|DOS|E|A|O)\b", " ", base)
    base = re.sub(r"\s+", " ", base).strip()
    return base


def clinic_slug(value):
    base = slugify(value)
    base = re.sub(r"\b(UNIDADE|MASCULINA|FEMININA|ACOLHIMENTO|ASSISTENCIA SOCIAL)\b", " ", base)
    base = re.sub(r"\s+", " ", base).strip()
    return base


def infer_city(*values):
    haystack = " ".join(slugify(value) for value in values if clean_text(value))
    for alias, city in CITY_ALIASES.items():
        if alias in haystack:
            return city
    return ""


def normalize_phone(value):
    digits = re.sub(r"\D", "", clean_text(value))
    return digits[:11] if digits else ""


def normalize_status(value):
    text = slugify(value)
    if not text:
        return "Ativo"
    if any(token in text for token in ["CANCEL", "INATIV", "REFORMA", "PAUSAD", "ENCERR"]):
        return "Inativo"
    return "Ativo"


def normalize_atendimento_status(value):
    text = slugify(value)
    if any(token in text for token in ["CANCEL", "INATIV", "REFORMA", "PAUSAD", "ENCERR"]):
        return "Cancelado"
    return "Agendado"


def classify_tipo_local(name, address="", sector=""):
    haystack = " ".join([slugify(name), slugify(address), slugify(sector)])
    if "CADEIA" in haystack or "RESSOCIALIZACAO" in haystack or "SOCIOEDUCATIVO" in haystack:
        return "Outro"
    if any(token in haystack for token in ["HOSPITAL", "UPA", "PSI", "PRONTO SOCORRO"]):
        return "Hospital"
    if any(token in haystack for token in ["UNIVERSIDADE", "FACULDADE", "INSTITUTO"]):
        return "Universidade"
    if any(token in haystack for token in ["CASA", "LAR", "ABRIGO", "ALBERGUE", "ACOLHIMENTO"]):
        return "Casa de Apoio"
    return "Clinica"


def infer_periodicidade(value):
    text = slugify(value)
    if not text:
        return ""
    if "QUINZEN" in text or "CADA 15" in text or "2 E 4" in text or "ULTIMA" in text:
        return "Quinzenal"
    if "TODO" in text or "TODOS" in text or "SEMANA" in text or "QUARTA E SABADO" in text:
        return "Semanal"
    if any(token in text for token in ["1 ", "2 ", "3 ", "4 ", "5 ", "PRIMEIR", "SEGUND", "TERCEIR", "QUART"]):
        return "Mensal"
    if "MES" in text:
        return "Mensal"
    return "Sob Escala"


def sql_quote(value):
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, (list, dict)):
        text = json.dumps(value, ensure_ascii=False)
    else:
        text = str(value)
    return "'" + text.replace("'", "''") + "'"


def parse_possible_date(value, month=None, year=None):
    if value in (None, "", "-"):
        return []

    if isinstance(value, datetime):
        if value.year == 1900 and month and year:
            return [date(year, month, value.day)]
        return [value.date()]

    if isinstance(value, date):
        if value.year == 1900 and month and year:
            return [date(year, month, value.day)]
        return [value]

    text = clean_text(value)
    if not text or text == "-":
        return []

    iso_match = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?: .*)?", text)
    if iso_match:
        return [date(int(iso_match.group(1)), int(iso_match.group(2)), int(iso_match.group(3)))]

    range_match = re.fullmatch(r"(\d{1,2})-(\d{1,2})/(\d{2})", text)
    if range_match and year:
        parsed_month = int(range_match.group(3))
        return [
            date(year, parsed_month, int(range_match.group(1))),
            date(year, parsed_month, int(range_match.group(2))),
        ]

    if month and year:
        days = [int(piece) for piece in re.findall(r"\b(\d{1,2})\b", text) if 1 <= int(piece) <= 31]
        if days:
            unique_days = []
            for day_value in days:
                if day_value not in unique_days:
                    unique_days.append(day_value)
            return [date(year, month, day_value) for day_value in unique_days]

    return []


def split_unidades(value):
    text = clean_text(value)
    if not text:
        return []
    pieces = re.split(r"\s*/\s*|\s*,\s*|\s+\+\s+", text)
    result = []
    for piece in pieces:
        cleaned = piece.strip()
        if cleaned:
            result.append(cleaned)
    return result


def parse_clinicas(ws):
    records = []
    current_sector = ""

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nome = clean_text(row[0] if len(row) > 0 else "")
        endereco = clean_text(row[1] if len(row) > 1 else "")
        cep = clean_text(row[2] if len(row) > 2 else "")
        periodicidade_raw = clean_text(row[3] if len(row) > 3 else "")
        horario = clean_text(row[4] if len(row) > 4 else "")
        responsaveis = clean_text(row[5] if len(row) > 5 else "")
        telefone = clean_text(row[6] if len(row) > 6 else "")
        status_raw = clean_text(row[7] if len(row) > 7 else "")

        if nome.startswith("SETOR ") and not endereco:
            current_sector = nome
            continue

        if not nome:
            continue

        records.append(
            {
                "source_row": row_index,
                "nome_local": nome,
                "tipo_local": classify_tipo_local(nome, endereco, current_sector),
                "cidade": infer_city(nome, endereco, current_sector),
                "endereco": endereco,
                "responsavel_local": responsaveis,
                "telefone_contato": normalize_phone(telefone),
                "periodicidade_preferencial": infer_periodicidade(periodicidade_raw),
                "status": normalize_status(status_raw),
                "raw_slug": clinic_slug(nome),
                "observacoes": " | ".join(
                    [
                        piece
                        for piece in [
                            f"Setor: {current_sector}" if current_sector else "",
                            f"Periodicidade original: {periodicidade_raw}" if periodicidade_raw else "",
                            f"Horario original: {horario}" if horario else "",
                            f"CEP: {cep}" if cep else "",
                            f"Origem: Clinica Detalhada linha {row_index}",
                            IMPORT_MARKER,
                        ]
                        if piece
                    ]
                ),
            }
        )

    return records


def parse_colaboradores(ws):
    records = []

    for row_index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nome = clean_text(row[0] if len(row) > 0 else "")
        unidade = clean_text(row[1] if len(row) > 1 else "")
        comum = clean_text(row[2] if len(row) > 2 else "")
        cidade = clean_text(row[3] if len(row) > 3 else "")
        ocupacao = clean_text(row[4] if len(row) > 4 else "")
        instrumento = clean_text(row[5] if len(row) > 5 else "")
        situacao = clean_text(row[6] if len(row) > 6 else "")
        cadastro = clean_text(row[7] if len(row) > 7 else "")
        telefone = clean_text(row[8] if len(row) > 8 else "")
        cpf = clean_text(row[9] if len(row) > 9 else "")

        if not nome:
            continue

        status = "Ativo" if "ATIVO" in slugify(situacao or "ATIVO") else "Inativo"
        observacoes = [
            f"Unidade(s): {unidade}" if unidade else "",
            f"Ocupacao: {ocupacao}" if ocupacao else "",
            f"Cadastro: {cadastro}" if cadastro else "",
            f"CPF: {cpf}" if cpf else "",
            f"Origem: Colaborador Atendente linha {row_index}",
            IMPORT_MARKER,
        ]

        records.append(
            {
                "source_row": row_index,
                "nome_completo": nome,
                "data_nascimento": None,
                "celular": normalize_phone(telefone),
                "email": "",
                "comum_congregacao": comum,
                "cidade": infer_city(cidade, comum, unidade) or clean_text(cidade),
                "instrumento": instrumento,
                "nivel": "",
                "disponibilidade": "",
                "apto_atendimentos": "Sim" if status == "Ativo" else "Nao",
                "status": status,
                "unidade_original": unidade,
                "unit_slugs": [clinic_slug(piece) for piece in split_unidades(unidade)],
                "observacoes": " | ".join(piece for piece in observacoes if piece),
            }
        )

    return records


def parse_acordo(ws):
    mapping = {}

    for row in ws.iter_rows(min_row=3, values_only=True):
        clinica = clean_text(row[0] if len(row) > 0 else "")
        dia_hora = clean_text(row[1] if len(row) > 1 else "")
        responsavel = clean_text(row[2] if len(row) > 2 else "")
        if not clinica:
            continue
        mapping[clinic_slug(clinica)] = {
            "dia_hora": dia_hora,
            "responsavel": responsavel,
        }

    return mapping


def parse_relacao_geral(ws):
    blocks = [(1, 2, 3, 4, 5), (9, 10, 11, 12, 13)]
    mapping = {}

    for row in ws.iter_rows(min_row=4, values_only=True):
        for cols in blocks:
            instituicao = clean_text(row[cols[0]] if len(row) > cols[0] else "")
            regiao = clean_text(row[cols[1]] if len(row) > cols[1] else "")
            reunioes = clean_text(row[cols[2]] if len(row) > cols[2] else "")
            media = clean_text(row[cols[3]] if len(row) > cols[3] else "")
            participantes = clean_text(row[cols[4]] if len(row) > cols[4] else "")
            if not instituicao:
                continue
            mapping[clinic_slug(instituicao)] = {
                "regiao": regiao,
                "reunioes_mes": reunioes,
                "media_participantes": media,
                "participantes_darpe": participantes,
            }

    return mapping


def extract_year_from_agenda_title(value):
    match = re.search(r"\b(20\d{2})\b", clean_text(value))
    return int(match.group(1)) if match else datetime.now().year


def build_clinic_lookup(clinicas):
    lookup = {}
    for record in clinicas:
        lookup[record["raw_slug"]] = record

    return lookup


def strip_time_suffix(name):
    text = clean_text(name)
    text = re.sub(r"\s*\((?:\d{1,2}h(?:\d{2})?|\d{1,2}:\d{2})\)\s*$", "", text, flags=re.IGNORECASE)
    return text.strip()


def match_clinica(name, clinic_lookup):
    slug = clinic_slug(name)
    if slug in clinic_lookup:
        return clinic_lookup[slug]

    candidates = []
    for candidate_slug, candidate in clinic_lookup.items():
        if slug and (slug in candidate_slug or candidate_slug in slug):
            candidates.append(candidate)

    if len(candidates) == 1:
        return candidates[0]

    return None


def parse_agenda(ws, clinic_lookup, musicais_por_local):
    rows = list(ws.iter_rows(values_only=True))
    title_year = extract_year_from_agenda_title(rows[0][0] if rows and rows[0] else "")
    month_headers = {}
    header_row = rows[1]

    for column_index in range(2, 8):
        month_name = slugify(header_row[column_index] if len(header_row) > column_index else "")
        if month_name in MONTHS:
            month_headers[column_index] = MONTHS[month_name]

    atendimentos = []
    unmatched = []

    for row_index, row in enumerate(rows[2:], start=3):
        local_nome = clean_text(row[0] if len(row) > 0 else "")
        dia_semana = clean_text(row[1] if len(row) > 1 else "")
        status_raw = clean_text(row[8] if len(row) > 8 else "")

        if not local_nome:
            continue

        clinica = match_clinica(local_nome, clinic_lookup)

        if not clinica:
            unmatched.append({"row": row_index, "local_nome": local_nome})

        for column_index, month in month_headers.items():
            raw_value = row[column_index] if len(row) > column_index else None
            for parsed_date in parse_possible_date(raw_value, month=month, year=title_year):
                matched_clinica = clinica or {
                    "id": None,
                    "nome_local": local_nome,
                    "tipo_local": classify_tipo_local(local_nome),
                    "cidade": "",
                    "responsavel_local": "",
                }
                musicos = musicais_por_local.get(matched_clinica.get("raw_slug", clinic_slug(local_nome)), [])
                atendimentos.append(
                    {
                        "data_atendimento": parsed_date.isoformat(),
                        "periodicidade": infer_periodicidade(dia_semana),
                        "local_id": matched_clinica.get("id"),
                        "local_nome": matched_clinica.get("nome_local") or local_nome,
                        "tipo_local": matched_clinica.get("tipo_local") or classify_tipo_local(local_nome),
                        "cidade": matched_clinica.get("cidade") or infer_city(local_nome),
                        "responsavel_ministerio": matched_clinica.get("responsavel_local", ""),
                        "musicos_ids": [musico["id"] for musico in musicos],
                        "musicos_nomes": ", ".join(musico["nome_completo"] for musico in musicos),
                        "quantidade_musicos": len(musicos),
                        "repertorio": "",
                        "observacoes": " | ".join(
                            piece
                            for piece in [
                                f"Agenda Simplificada linha {row_index}",
                                f"Dia semana original: {dia_semana}" if dia_semana else "",
                                f"Status agenda original: {status_raw}" if status_raw else "",
                                IMPORT_MARKER,
                            ]
                            if piece
                        ),
                        "status": normalize_atendimento_status(status_raw),
                        "proxima_visita": None,
                    }
                )

    return atendimentos, unmatched


def enrich_clinicas(clinicas, acordo_map, relacao_map):
    for record in clinicas:
        acordo = acordo_map.get(record["raw_slug"])
        relacao = relacao_map.get(record["raw_slug"])
        extras = []

        if acordo:
            if not record["responsavel_local"]:
                record["responsavel_local"] = acordo["responsavel"]
            if not record["periodicidade_preferencial"]:
                record["periodicidade_preferencial"] = infer_periodicidade(acordo["dia_hora"])
            extras.append(f"Acordo 2103: {acordo['dia_hora']}")
            if acordo["responsavel"]:
                extras.append(f"Responsavel ministerio: {acordo['responsavel']}")

        if relacao:
            extras.append(
                "Relacao Geral: "
                + ", ".join(
                    piece
                    for piece in [
                        f"regiao {relacao['regiao']}" if relacao["regiao"] else "",
                        f"reunioes/mes {relacao['reunioes_mes']}" if relacao["reunioes_mes"] else "",
                        f"media participantes {relacao['media_participantes']}" if relacao["media_participantes"] else "",
                        f"participantes DARPE {relacao['participantes_darpe']}" if relacao["participantes_darpe"] else "",
                    ]
                    if piece
                )
            )

        if extras:
            record["observacoes"] = record["observacoes"] + " | " + " | ".join(extras)


def link_colaboradores_to_clinicas(colaboradores, clinicas):
    clinic_lookup = build_clinic_lookup(clinicas)
    musicais_por_local = defaultdict(list)

    for colaborador in colaboradores:
        seen = set()
        for unit_slug in colaborador["unit_slugs"]:
            if not unit_slug:
                continue
            clinica = clinic_lookup.get(unit_slug)
            if not clinica:
                clinica = match_clinica(unit_slug, clinic_lookup)
            if not clinica:
                continue
            clinica_slug = clinica["raw_slug"]
            if clinica_slug in seen:
                continue
            seen.add(clinica_slug)
            musicais_por_local[clinica_slug].append(colaborador)

    return musicais_por_local


def assign_ids(records, start_id):
    for offset, record in enumerate(records):
        record["id"] = start_id + offset


def to_seed_sql(clinicas, musicos, atendimentos):
    lines = [
        "-- Seed gerado automaticamente por scripts/darpe_import.py",
        f"-- Marcador de importacao: {IMPORT_MARKER}",
        "BEGIN;",
        f"DELETE FROM public.darpe_atendimentos WHERE observacoes LIKE '%{IMPORT_MARKER}%';",
        f"DELETE FROM public.darpe_clinicas WHERE observacoes LIKE '%{IMPORT_MARKER}%';",
        f"DELETE FROM public.darpe_musicos WHERE observacoes LIKE '%{IMPORT_MARKER}%';",
        "",
    ]

    for record in musicos:
        lines.append(
            "INSERT INTO public.darpe_musicos "
            "(id, nome_completo, data_nascimento, celular, email, comum_congregacao, cidade, instrumento, nivel, disponibilidade, apto_atendimentos, observacoes, status) "
            "OVERRIDING SYSTEM VALUE VALUES "
            f"({sql_quote(record['id'])}, {sql_quote(record['nome_completo'])}, {sql_quote(record['data_nascimento'])}, "
            f"{sql_quote(record['celular'])}, {sql_quote(record['email'])}, {sql_quote(record['comum_congregacao'])}, "
            f"{sql_quote(record['cidade'])}, {sql_quote(record['instrumento'])}, {sql_quote(record['nivel'])}, "
            f"{sql_quote(record['disponibilidade'])}, {sql_quote(record['apto_atendimentos'])}, {sql_quote(record['observacoes'])}, "
            f"{sql_quote(record['status'])});"
        )

    lines.append("")

    for record in clinicas:
        lines.append(
            "INSERT INTO public.darpe_clinicas "
            "(id, nome_local, tipo_local, cidade, endereco, responsavel_local, telefone_contato, periodicidade_preferencial, observacoes, status) "
            "OVERRIDING SYSTEM VALUE VALUES "
            f"({sql_quote(record['id'])}, {sql_quote(record['nome_local'])}, {sql_quote(record['tipo_local'])}, "
            f"{sql_quote(record['cidade'])}, {sql_quote(record['endereco'])}, {sql_quote(record['responsavel_local'])}, "
            f"{sql_quote(record['telefone_contato'])}, {sql_quote(record['periodicidade_preferencial'])}, {sql_quote(record['observacoes'])}, "
            f"{sql_quote(record['status'])});"
        )

    lines.append("")

    for record in atendimentos:
        lines.append(
            "INSERT INTO public.darpe_atendimentos "
            "(data_atendimento, periodicidade, local_id, local_nome, tipo_local, cidade, responsavel_ministerio, musicos_ids, musicos_nomes, quantidade_musicos, repertorio, observacoes, status, proxima_visita) "
            "VALUES "
            f"({sql_quote(record['data_atendimento'])}, {sql_quote(record['periodicidade'])}, {sql_quote(record['local_id'])}, "
            f"{sql_quote(record['local_nome'])}, {sql_quote(record['tipo_local'])}, {sql_quote(record['cidade'])}, "
            f"{sql_quote(record['responsavel_ministerio'])}, {sql_quote(record['musicos_ids'])}::jsonb, {sql_quote(record['musicos_nomes'])}, "
            f"{sql_quote(record['quantidade_musicos'])}, {sql_quote(record['repertorio'])}, {sql_quote(record['observacoes'])}, "
            f"{sql_quote(record['status'])}, {sql_quote(record['proxima_visita'])});"
        )

    lines.extend(
        [
            "",
            "COMMIT;",
            "",
            "-- Se necessario, ajuste a sequencia apos a importacao:",
            "-- SELECT setval('public.darpe_musicos_id_seq', GREATEST((SELECT COALESCE(MAX(id), 1) FROM public.darpe_musicos), nextval('public.darpe_musicos_id_seq')));",
            "-- SELECT setval('public.darpe_clinicas_id_seq', GREATEST((SELECT COALESCE(MAX(id), 1) FROM public.darpe_clinicas), nextval('public.darpe_clinicas_id_seq')));",
        ]
    )

    return "\n".join(lines) + "\n"


def build_report(summary):
    lines = [
        "# Importacao DARPE",
        "",
        f"- Arquivo analisado: `{WORKBOOK_PATH.as_posix()}`",
        f"- Clinicas/localidades extraidas: **{summary['counts']['clinicas']}**",
        f"- Colaboradores/musicos extraidos: **{summary['counts']['musicos']}**",
        f"- Atendimentos gerados a partir da agenda: **{summary['counts']['atendimentos']}**",
        f"- Locais da agenda sem correspondencia automatica: **{summary['counts']['agenda_nao_correspondida']}**",
        f"- Marcador tecnico: `{IMPORT_MARKER}`",
        "",
        "## Arquivos gerados",
        "",
        "- `scripts/darpe_import_output/darpe_seed.sql`",
        "- `scripts/darpe_import_output/darpe_import_summary.json`",
        "- `scripts/darpe_import_output/darpe_preview.json`",
        "",
        "## Observacoes do mapeamento",
        "",
        "- `Colaborador Atendente` alimenta `darpe_musicos`.",
        "- `Clinica Detalhada` alimenta `darpe_clinicas`.",
        "- `Agenda Simplificada` foi explodida em registros individuais para `darpe_atendimentos`.",
        "- `Acordo 2103` e `Relacao Geral` foram usados para enriquecer responsaveis e observacoes dos locais.",
    ]

    if summary["agenda_unmatched"]:
        lines.extend(
            [
                "",
                "## Locais da agenda sem match automatico",
                "",
            ]
        )
        for item in summary["agenda_unmatched"][:20]:
            lines.append(f"- Linha {item['row']}: {item['local_nome']}")

    return "\n".join(lines) + "\n"


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
    agenda_ws = wb["Agenda Simplificada"]
    clinicas_ws = wb["Clínica Detalhada"]
    colaboradores_ws = wb["Colaborador Atendente"]
    relacao_ws = wb["Relação Geral"]
    acordo_ws = wb["Acordo 2103"]

    clinicas = parse_clinicas(clinicas_ws)
    colaboradores = parse_colaboradores(colaboradores_ws)
    acordo_map = parse_acordo(acordo_ws)
    relacao_map = parse_relacao_geral(relacao_ws)

    enrich_clinicas(clinicas, acordo_map, relacao_map)
    clinic_lookup = build_clinic_lookup(clinicas)
    _, unmatched = parse_agenda(agenda_ws, clinic_lookup, {})

    if unmatched:
        seen_unmatched = set()
        for item in unmatched:
            local_nome = item["local_nome"]
            stripped_name = strip_time_suffix(local_nome)
            slug = clinic_slug(stripped_name)
            if not slug or slug in clinic_lookup or slug in seen_unmatched:
                continue
            seen_unmatched.add(slug)
            acordo = acordo_map.get(slug)
            relacao = relacao_map.get(slug)
            synthetic = {
                "source_row": item["row"],
                "nome_local": stripped_name,
                "tipo_local": classify_tipo_local(stripped_name),
                "cidade": infer_city(stripped_name, relacao["regiao"] if relacao else ""),
                "endereco": "",
                "responsavel_local": acordo["responsavel"] if acordo else "",
                "telefone_contato": "",
                "periodicidade_preferencial": infer_periodicidade(acordo["dia_hora"]) if acordo else "",
                "status": "Ativo",
                "raw_slug": slug,
                "observacoes": " | ".join(
                    piece
                    for piece in [
                        f"Clinica sintetica criada a partir da agenda linha {item['row']}",
                        f"Acordo 2103: {acordo['dia_hora']}" if acordo else "",
                        (
                            "Relacao Geral: "
                            + ", ".join(
                                part
                                for part in [
                                    f"regiao {relacao['regiao']}" if relacao and relacao["regiao"] else "",
                                    f"reunioes/mes {relacao['reunioes_mes']}" if relacao and relacao["reunioes_mes"] else "",
                                    f"media participantes {relacao['media_participantes']}" if relacao and relacao["media_participantes"] else "",
                                    f"participantes DARPE {relacao['participantes_darpe']}" if relacao and relacao["participantes_darpe"] else "",
                                ]
                                if part
                            )
                        ) if relacao else "",
                        IMPORT_MARKER,
                    ]
                    if piece
                ),
            }
            clinicas.append(synthetic)

    assign_ids(clinicas, 910000)
    assign_ids(colaboradores, 920000)

    musicais_por_local = link_colaboradores_to_clinicas(colaboradores, clinicas)
    clinic_lookup = build_clinic_lookup(clinicas)
    atendimentos, unmatched = parse_agenda(agenda_ws, clinic_lookup, musicais_por_local)

    summary = {
        "counts": {
            "clinicas": len(clinicas),
            "musicos": len(colaboradores),
            "atendimentos": len(atendimentos),
            "agenda_nao_correspondida": len(unmatched),
        },
        "agenda_unmatched": unmatched,
    }

    preview = {
        "clinicas": clinicas[:8],
        "musicos": colaboradores[:8],
        "atendimentos": atendimentos[:12],
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "darpe_seed.sql").write_text(to_seed_sql(clinicas, colaboradores, atendimentos), encoding="utf-8")
    (OUTPUT_DIR / "darpe_import_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "darpe_preview.json").write_text(json.dumps(preview, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUT_DIR / "README.md").write_text(build_report(summary), encoding="utf-8")

    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
